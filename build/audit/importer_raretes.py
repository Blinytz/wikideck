#!/usr/bin/env python3
"""Reprend le tableau raretes_pv.xlsx retouche par l'utilisateur.

Toute rarete ou tout PV qui differe du jeu est applique et marque manuel
(rareteManuel / pvManuel) : le prochain recalcul automatique le respectera.

Usage : python importer_raretes.py [chemin.xlsx]
"""
import sys, json
from pathlib import Path
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8')
ICI = Path(__file__).resolve().parent
RACINE = ICI.parent.parent
CODES = {'Commune': 'commune', 'Rare': 'rare', 'Épique': 'epique', 'Mythique': 'mythique',
         'Légendaire': 'legendaire', 'Panthéon': 'pantheon'}


def main():
    f = Path(sys.argv[1]) if len(sys.argv) > 1 else ICI / 'raretes_pv.xlsx'
    ws = load_workbook(f, data_only=True)['Cartes']
    ent = [c.value for c in ws[1]]
    i_id, i_r, i_pv = ent.index('id'), ent.index('Rareté'), ent.index('PV')
    voulu = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[i_id]:
            voulu[row[i_id]] = (CODES.get(str(row[i_r]).strip(), None), row[i_pv])
    idx = json.loads((RACINE / 'data' / 'collections.json').read_text(encoding='utf-8'))
    nr = npv = 0
    for e in idx['collections']:
        p = RACINE / e['fichier']
        d = json.loads(p.read_text(encoding='utf-8'))
        for c in d['cartes']:
            v = voulu.get(c['id'])
            if not v:
                continue
            r, pv = v
            if r and r != c['rarete']:
                c['rarete'], c['rareteManuel'] = r, True; nr += 1
            if pv is not None and int(pv) != c.get('pv'):
                c['pv'], c['pvManuel'] = int(pv), True; npv += 1
        p.write_text(json.dumps(d, ensure_ascii=False), encoding='utf-8')
    print(f'{nr} rarete(s) et {npv} PV modifies')


if __name__ == '__main__':
    main()
